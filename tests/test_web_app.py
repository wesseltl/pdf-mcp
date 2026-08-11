"""Tests for the no-code local browser app."""
from __future__ import annotations

import http.client
import json
import threading
import unittest
from io import BytesIO
from urllib.parse import quote

from docx import Document
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from pdf_mcp.web_app import create_server


def pdf_with_table() -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4)
    table = Table([["Item", "Qty"], ["Widget", "3"], ["Bolt", "10"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    document.build([table])
    return buffer.getvalue()


def pdf_without_table() -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4)
    document.build([Paragraph("A document without a table", getSampleStyleSheet()["BodyText"])])
    return buffer.getvalue()


def docx_with_table() -> bytes:
    buffer = BytesIO()
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Analyte"
    table.rows[0].cells[1].text = "Result"
    cells = table.add_row().cells
    cells[0].text = "pH"
    cells[1].text = "7.2"
    document.save(buffer)
    return buffer.getvalue()


class TestLocalWebApp(unittest.TestCase):
    def setUp(self):
        self.server, self.state = create_server(0)
        self.port = self.server.server_address[1]
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.state.clear()
        self.thread.join(timeout=2)

    def request(self, method, path, body=None, headers=None):
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=10)
        connection.request(method, path, body=body, headers=headers or {})
        response = connection.getresponse()
        content = response.read()
        result = (response.status, dict(response.getheaders()), content)
        connection.close()
        return result

    def conversion_headers(self, content: bytes) -> dict[str, str]:
        return {
            "Content-Type": "application/octet-stream",
            "Content-Length": str(len(content)),
            "X-Pdf-Mcp-Session": self.state.session_token,
        }

    def test_home_page_injects_session_and_security_headers(self):
        status, headers, content = self.request("GET", "/")

        self.assertEqual(status, 200)
        self.assertIn(self.state.session_token.encode(), content)
        self.assertNotIn(b"__PDF_MCP_SESSION__", content)
        self.assertIn("default-src 'self'", headers["Content-Security-Policy"])
        self.assertEqual(headers["X-Frame-Options"], "DENY")
        self.assertNotIn("Access-Control-Allow-Origin", headers)

    def test_conversion_requires_the_browser_session_token(self):
        content = pdf_with_table()
        status, _headers, response = self.request(
            "POST",
            "/api/convert?filename=invoice.pdf&format=xlsx",
            body=content,
            headers={"Content-Length": str(len(content))},
        )

        self.assertEqual(status, 403)
        self.assertIn("not authorized", json.loads(response)["error"])

    def test_cross_origin_conversion_is_rejected(self):
        content = pdf_with_table()
        headers = self.conversion_headers(content)
        headers["Origin"] = "https://malicious.example"
        status, _headers, response = self.request(
            "POST",
            "/api/convert?filename=invoice.pdf&format=xlsx",
            body=content,
            headers=headers,
        )

        self.assertEqual(status, 403)
        self.assertIn("local requests", json.loads(response)["error"])

    def test_pdf_conversion_returns_preview_and_downloadable_workbook(self):
        content = pdf_with_table()
        status, _headers, response = self.request(
            "POST",
            "/api/convert?filename=invoice.pdf&format=xlsx&merge_multipage=1",
            body=content,
            headers=self.conversion_headers(content),
        )

        self.assertEqual(status, 200)
        payload = json.loads(response)
        self.assertEqual(payload["n_tables"], 1)
        self.assertEqual(payload["tables"][0]["rows"][0], ["Item", "Qty"])
        self.assertEqual(payload["tables"][0]["rows"][1], ["Widget", "3"])
        self.assertEqual(payload["output_name"], "invoice-tables.xlsx")

        download_status, download_headers, workbook_content = self.request(
            "GET",
            f"/api/download/{payload['download_id']}",
            headers={"X-Pdf-Mcp-Session": self.state.session_token},
        )
        self.assertEqual(download_status, 200)
        self.assertIn("invoice-tables.xlsx", download_headers["Content-Disposition"])
        workbook = load_workbook(BytesIO(workbook_content))
        self.assertEqual(workbook["Table 1"]["B1"].value, "invoice.pdf")
        self.assertEqual(workbook["Table 1"]["A7"].value, "Widget")

    def test_original_filename_is_reduced_to_a_safe_display_name(self):
        content = pdf_with_table()
        filename = quote("../../customer invoice.pdf")
        status, _headers, response = self.request(
            "POST",
            f"/api/convert?filename={filename}&format=json",
            body=content,
            headers=self.conversion_headers(content),
        )

        self.assertEqual(status, 200)
        payload = json.loads(response)
        self.assertEqual(payload["input_name"], "customer invoice.pdf")
        self.assertEqual(payload["output_name"], "customer invoice-tables.json")

    def test_word_document_can_be_downloaded_as_csv_or_json(self):
        content = docx_with_table()
        for output_type, expected in (("csv", b"Analyte,Result"), ("json", b'"source_type": "docx"')):
            with self.subTest(output_type=output_type):
                status, _headers, response = self.request(
                    "POST",
                    f"/api/convert?filename=results.docx&format={output_type}",
                    body=content,
                    headers=self.conversion_headers(content),
                )
                self.assertEqual(status, 200)
                payload = json.loads(response)
                self.assertEqual(payload["tables"][0]["rows"][1], ["pH", "7.2"])
                download_status, _download_headers, download = self.request(
                    "GET",
                    f"/api/download/{payload['download_id']}",
                    headers={"X-Pdf-Mcp-Session": self.state.session_token},
                )
                self.assertEqual(download_status, 200)
                self.assertIn(expected, download)

    def test_no_table_response_explains_the_ocr_boundary(self):
        content = pdf_without_table()
        status, _headers, response = self.request(
            "POST",
            "/api/convert?filename=scan.pdf&format=xlsx",
            body=content,
            headers=self.conversion_headers(content),
        )

        self.assertEqual(status, 422)
        payload = json.loads(response)
        self.assertIn("No tables", payload["error"])
        self.assertIn("OCR", payload["hint"])

    def test_extension_and_signature_are_validated(self):
        content = b"not a pdf"
        status, _headers, response = self.request(
            "POST",
            "/api/convert?filename=notes.txt&format=xlsx",
            body=content,
            headers=self.conversion_headers(content),
        )
        self.assertEqual(status, 400)
        self.assertIn("PDF and Word", json.loads(response)["error"])

        status, _headers, response = self.request(
            "POST",
            "/api/convert?filename=notes.pdf&format=xlsx",
            body=content,
            headers=self.conversion_headers(content),
        )
        self.assertEqual(status, 400)
        self.assertIn("does not appear", json.loads(response)["error"])

    def test_download_requires_the_same_browser_session(self):
        status, _headers, response = self.request("GET", "/api/download/unknown")
        self.assertEqual(status, 403)
        self.assertIn("not authorized", json.loads(response)["error"])

    def test_shutdown_requires_the_browser_session(self):
        status, _headers, _response = self.request("POST", "/api/shutdown", body=b"")
        self.assertEqual(status, 403)


if __name__ == "__main__":
    unittest.main()
